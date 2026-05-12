"""Convertit PRESENTATION_SONAR_CLOUD_GITHUB_ACTIONS.md en .docx (python-docx)."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "PRESENTATION_SONAR_CLOUD_GITHUB_ACTIONS.md"
OUT = ROOT / "PRESENTATION_SONAR_CLOUD_GITHUB_ACTIONS.docx"


def _add_segments_with_bold(p, text: str) -> None:
    parts = text.split("**")
    for idx, chunk in enumerate(parts):
        if not chunk:
            continue
        run = p.add_run(chunk)
        if idx % 2 == 1:
            run.bold = True


def flush_paragraphs(doc: Document, buf: list[str]) -> None:
    raw = "\n".join(buf).strip()
    if not raw:
        return
    for block in raw.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_segments_with_bold(p, block[2:])
            continue
        p = doc.add_paragraph()
        _add_segments_with_bold(p, block)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    def clean_cell(cell: str) -> str:
        return cell.replace("**", "").strip()

    if len(rows) >= 2:
        sep = "".join(rows[1]).replace("|", "").replace("-", "").replace(":", "").strip()
        if not sep:
            rows = [rows[0]] + rows[2:]
    rows = [[clean_cell(c) for c in r] for r in rows]
    cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    for ri, r in enumerate(rows):
        for ci in range(cols):
            tbl.rows[ri].cells[ci].text = r[ci] if ci < len(r) else ""


def shading_code(paragraph):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    paragraph._p.get_or_add_pPr().append(shading)


def flush_code_block(doc: Document, lines: list[str]) -> None:
    if not lines:
        return
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    shading_code(p)


def build() -> Document:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.core_properties.title = "GitHub et SonarCloud — projet NAFISSA"
    doc.core_properties.subject = "Présentation équipe CI / qualité"

    buf: list[str] = []
    i = 0
    code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if s.startswith("```"):
            flush_paragraphs(doc, buf)
            buf.clear()
            if code:
                flush_code_block(doc, code_buf)
                code_buf.clear()
                code = False
            else:
                code = True
            i += 1
            continue

        if code:
            code_buf.append(line)
            i += 1
            continue

        if s == "---":
            flush_paragraphs(doc, buf)
            buf.clear()
            i += 1
            continue

        if line.startswith("# ") and not line.startswith("## "):
            flush_paragraphs(doc, buf)
            buf.clear()
            doc.add_heading(line[2:].strip(), 0)
            i += 1
            continue
        if line.startswith("## "):
            flush_paragraphs(doc, buf)
            buf.clear()
            doc.add_heading(line[3:].strip(), 1)
            i += 1
            continue
        if line.startswith("### "):
            flush_paragraphs(doc, buf)
            buf.clear()
            doc.add_heading(line[4:].strip(), 2)
            i += 1
            continue

        if s.startswith("|"):
            flush_paragraphs(doc, buf)
            buf.clear()
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                table_rows.append(row)
                i += 1
            add_table(doc, table_rows)
            continue

        buf.append(line)
        i += 1

    flush_paragraphs(doc, buf)
    return doc


def main() -> None:
    doc = build()
    candidates = [OUT, OUT.with_stem(f"{OUT.stem}_nouveau")]
    last_err = None
    for target in candidates:
        try:
            doc.save(target)
            print(f"OK: {target}")
            return
        except PermissionError as e:
            last_err = e
            continue
    raise last_err  # pragma: no cover


if __name__ == "__main__":
    main()
